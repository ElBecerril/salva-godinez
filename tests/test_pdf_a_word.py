"""Regresion de tools/pdf_tools.pdf_to_word_do (PDF a Word).

Antes solo copiaba el texto: se perdian imagenes, tablas y formato. Ahora usa
pdf2docx (modo "fiel") y, si no esta o falla con ese PDF, cae al modo "texto".
Se prueba:
  - modo fiel: el .docx conserva la imagen y la tabla del PDF
  - modo texto (respaldo): sigue convirtiendo el texto sin imagen ni tabla
  - si pdf2docx truena a la mitad, cae al modo texto en vez de fallar
  - PDF cifrado con clave y PDF corrupto: error limpio, no excepcion
  - un escaneo (solo imagen) no se reporta como "vacio" en modo fiel

Se omiten los casos del modo fiel si no esta instalado pdf2docx:

    python3 tests/test_pdf_a_word.py
"""

import os
import sys
import tempfile

_RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _RAIZ)

import fitz  # noqa: E402
import pypdf  # noqa: E402
from docx import Document  # noqa: E402
from PIL import Image  # noqa: E402

from tools import pdf_tools as P  # noqa: E402


fallos = 0


def check(nombre, cond, extra=""):
    global fallos
    fallos += not cond
    marca = "OK " if cond else "FAIL"
    detalle = f" -> {extra}" if not cond and extra else ""
    print(f"[{marca}] {nombre}{detalle}")


HAY_P2D = P._import_pdf2docx() is not None


def crear_pdf_completo(path, png):
    """Una pagina con titulo, parrafo, una imagen y una tabla de 4x3 con lineas."""
    d = fitz.open()
    p = d.new_page()
    p.insert_text((72, 80), "Informe de prueba", fontsize=20)
    p.insert_text((72, 110), "Parrafo normal para comprobar la conversion.", fontsize=11)
    p.insert_image(fitz.Rect(72, 130, 272, 230), filename=png)
    for r in range(4):
        for c in range(3):
            rect = fitz.Rect(72 + c * 120, 260 + r * 30, 72 + (c + 1) * 120, 260 + (r + 1) * 30)
            p.draw_rect(rect)
            p.insert_text((rect.x0 + 6, rect.y1 - 10), f"R{r}C{c}", fontsize=11)
    d.save(path)
    d.close()
    return path


def convertir(pdf, salida):
    return P.pdf_to_word_do(fitz, Document, pdf, salida, overwrite=True)


def resumen(docx):
    d = Document(docx)
    texto = " ".join(p.text for p in d.paragraphs) + " ".join(
        c.text for t in d.tables for row in t.rows for c in row.cells)
    return len(d.tables), len(d.inline_shapes), texto


with tempfile.TemporaryDirectory() as tmp:
    png = os.path.join(tmp, "rojo.png")
    Image.new("RGB", (300, 150), (200, 40, 40)).save(png)
    pdf = crear_pdf_completo(os.path.join(tmp, "completo.pdf"), png)

    if HAY_P2D:
        r = convertir(pdf, os.path.join(tmp, "fiel.docx"))
        check("modo fiel: convierte ok", r["ok"] and r["modo"] == "fiel", r)
        tablas, imagenes, texto = resumen(os.path.join(tmp, "fiel.docx"))
        check("modo fiel: conserva la tabla", tablas == 1, tablas)
        check("modo fiel: conserva la imagen", imagenes == 1, imagenes)
        check("modo fiel: conserva el texto de la tabla", "R3C2" in texto, texto[:80])
        check("modo fiel: no se reporta vacio", r["empty"] is False)
    else:
        print("[SKIP] pdf2docx no instalado: se omiten los casos del modo fiel")

    # Respaldo: sin pdf2docx se usa la extraccion de texto de siempre.
    original = P._import_pdf2docx
    P._import_pdf2docx = lambda: None
    try:
        r = convertir(pdf, os.path.join(tmp, "texto.docx"))
        check("modo texto: convierte ok", r["ok"] and r["modo"] == "texto", r)
        tablas, imagenes, texto = resumen(os.path.join(tmp, "texto.docx"))
        check("modo texto: conserva el texto", "Informe de prueba" in texto and "R3C2" in texto)
        check("modo texto: sin tabla ni imagen (limitacion conocida)", tablas == 0 and imagenes == 0)
    finally:
        P._import_pdf2docx = original

    # pdf2docx falla a la mitad: cae al modo texto en vez de dar error.
    class ConvertidorRoto:
        def __init__(self, path):
            pass

        def convert(self, output):
            raise RuntimeError("pdf2docx truena con este PDF")

        def close(self):
            pass

    P._import_pdf2docx = lambda: ConvertidorRoto
    try:
        r = convertir(pdf, os.path.join(tmp, "respaldo.docx"))
        check("pdf2docx falla: cae al modo texto", r["ok"] and r["modo"] == "texto", r)
        check("pdf2docx falla: el .docx existe", os.path.isfile(os.path.join(tmp, "respaldo.docx")))
    finally:
        P._import_pdf2docx = original

    # Escaneo: solo imagen, sin texto. En modo fiel no es "vacio".
    if HAY_P2D:
        d = fitz.open()
        d.new_page().insert_image(fitz.Rect(50, 50, 350, 250), filename=png)
        esc = os.path.join(tmp, "escaneo.pdf")
        d.save(esc)
        d.close()
        r = convertir(esc, os.path.join(tmp, "escaneo.docx"))
        check("escaneo: convierte y no se reporta vacio", r["ok"] and r["empty"] is False, r)

    # PDF sin texto ni imagenes: vacio.
    w = pypdf.PdfWriter()
    w.add_blank_page(200, 200)
    en_blanco = os.path.join(tmp, "blanco.pdf")
    with open(en_blanco, "wb") as f:
        w.write(f)
    r = convertir(en_blanco, os.path.join(tmp, "blanco.docx"))
    check("pagina en blanco: ok y se reporta vacia", r["ok"] and r["empty"] is True, r)

    # Errores limpios.
    malo = os.path.join(tmp, "malo.pdf")
    with open(malo, "wb") as f:
        f.write(b"esto no es un pdf")
    r = convertir(malo, os.path.join(tmp, "malo.docx"))
    check("PDF corrupto: error limpio", r["ok"] is False and r["error"] == "open_error", r)

    cif = os.path.join(tmp, "cifrado.pdf")
    w = pypdf.PdfWriter()
    w.add_blank_page(200, 200)
    w.encrypt(user_password="secreta", algorithm="AES-256")
    with open(cif, "wb") as f:
        w.write(f)
    r = convertir(cif, os.path.join(tmp, "cifrado.docx"))
    check("PDF con clave: error 'encrypted'", r["ok"] is False and r["error"] == "encrypted", r)

    # Sin overwrite no pisa un archivo existente.
    dest = os.path.join(tmp, "existente.docx")
    with open(dest, "w") as f:
        f.write("no tocar")
    r = P.pdf_to_word_do(fitz, Document, pdf, dest)
    check("sin overwrite: no pisa el existente",
          r["ok"] and r["output"] != dest and open(dest).read() == "no tocar", r)


print()
if fallos:
    print(f"RESULTADO: {fallos} FALLO(S)")
    sys.exit(1)
print("RESULTADO: TODO VERDE")
